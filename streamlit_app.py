
import io
import os
import zipfile
from pathlib import Path

import streamlit as st
from openpyxl import load_workbook

TEMPLATE_NAME = "!Шаблон заливки компенсаций.xlsx"
NUM_FMT = "#,##0.00"

SOURCE_TO_TARGET = {
    "Наименование акции": "Наименование акции",
    "механика для сложной или процент скидки для ценовой": "механика для сложной или процент скидки для ценовой",
    "код товара": "код товара",
    "Товар наименование": "Товар наименование",
    "ЦП": "ЦП",
    "ЦР": "ЦР",
    "ЦР Акционная": "ЦР Акционная",
    "продажи ИТОГО": "продажи ИТОГО ",
    "Акционные продажи (компенсируемые)": "Акционные продажи (компенсируемые)",
    "компенсации ": "КОМПЕНСАЦИЯ ПОСТАВЩИКА",
}

NUMERIC_TARGETS = [
    "ЦП", "ЦР", "ЦР Акционная",
    "продажи ИТОГО",
    "Акционные продажи (компенсируемые)",
    "компенсации ",
]

def norm(v):
    return "" if v is None else str(v).strip()

def is_zero(v):
    if v is None or str(v).strip() == "":
        return True
    try:
        return float(str(v).replace(" ", "").replace(",", ".")) == 0
    except Exception:
        return False

def find_header(ws):
    for r in range(1, min(ws.max_row, 40) + 1):
        vals = [norm(ws.cell(r, c).value) for c in range(1, ws.max_column + 1)]
        if "КОМПЕНСАЦИЯ ПОСТАВЩИКА" in vals and "код товара" in vals:
            return r
    raise ValueError("Не найдена строка заголовков исходного файла")

def make_output_name(filename):
    stem = Path(filename).stem
    if stem.startswith("Компенсации ИМ "):
        stem = "Шаблон заливки " + stem[len("Компенсации ИМ "):]
    elif stem.startswith("Компенсации "):
        stem = "Шаблон заливки " + stem[len("Компенсации "):]
    else:
        stem = "Шаблон заливки " + stem
    return stem + ".xlsx"

def process_file(uploaded_file, passport):
    swb = load_workbook(uploaded_file, data_only=True)
    sws = swb[swb.sheetnames[0]]

    template_path = Path(__file__).parent / TEMPLATE_NAME
    twb = load_workbook(template_path)
    tws = twb[twb.sheetnames[0]]

    hr = find_header(sws)
    src_headers = {norm(sws.cell(hr, c).value): c for c in range(1, sws.max_column + 1)}
    dst_headers = {norm(tws.cell(3, c).value): c for c in range(1, tws.max_column + 1)}

    missing_src = [norm(s) for s in SOURCE_TO_TARGET.values() if norm(s) not in src_headers]
    missing_dst = [norm(d) for d in SOURCE_TO_TARGET if norm(d) not in dst_headers]
    for required in ("Номер паспорта акций", "Направление бизнеса"):
        if required not in dst_headers:
            missing_dst.append(required)

    if missing_src or missing_dst:
        raise ValueError(f"Не хватает колонок. Исходник: {missing_src}; шаблон: {missing_dst}")

    dest_row = 4
    kept = skipped = 0
    comp_col = src_headers["КОМПЕНСАЦИЯ ПОСТАВЩИКА"]

    for r in range(hr + 1, sws.max_row + 1):
        if all(sws.cell(r, c).value is None for c in range(1, sws.max_column + 1)):
            continue

        comp = sws.cell(r, comp_col).value
        if is_zero(comp):
            skipped += 1
            continue

        for dst_name, src_name in SOURCE_TO_TARGET.items():
            dc = dst_headers[norm(dst_name)]
            sc = src_headers[norm(src_name)]
            tws.cell(dest_row, dc).value = sws.cell(r, sc).value

        tws.cell(dest_row, dst_headers["Номер паспорта акций"]).value = int(passport)
        tws.cell(dest_row, dst_headers["Направление бизнеса"]).value = 2
        kept += 1
        dest_row += 1

    for name in NUMERIC_TARGETS:
        col = dst_headers[norm(name)]
        for r in range(4, dest_row):
            tws.cell(r, col).number_format = NUM_FMT

    output = io.BytesIO()
    twb.save(output)
    output.seek(0)
    return output.getvalue(), kept, skipped

st.set_page_config(page_title="Компенсации → шаблоны", page_icon="📄", layout="centered")

st.title("Компенсации → шаблоны")
st.caption("Массовое формирование файлов для рабочей утилиты")

passport = st.text_input("Паспорт акции", value="151510")
files = st.file_uploader(
    "Загрузи файлы «Компенсации … .xlsx»",
    type=["xlsx"],
    accept_multiple_files=True,
)

st.info("Направление бизнеса всегда 2. Строки с компенсацией поставщика = 0 удаляются.")

if st.button("Сформировать шаблоны", type="primary", use_container_width=True):
    if not passport.strip().isdigit():
        st.error("Паспорт акции должен быть числом.")
    elif not files:
        st.error("Сначала загрузи хотя бы один Excel-файл.")
    else:
        zip_buffer = io.BytesIO()
        ok = bad = total_rows = total_skipped = 0
        results = []

        with zipfile.ZipFile(zip_buffer, "w", zipfile.ZIP_DEFLATED) as zf:
            for uploaded in files:
                try:
                    uploaded.seek(0)
                    data, kept, skipped = process_file(uploaded, passport.strip())
                    out_name = make_output_name(uploaded.name)
                    zf.writestr(out_name, data)
                    results.append(("✅", uploaded.name, kept, skipped, ""))
                    ok += 1
                    total_rows += kept
                    total_skipped += skipped
                except Exception as e:
                    results.append(("❌", uploaded.name, 0, 0, str(e)))
                    bad += 1

        zip_buffer.seek(0)

        if ok:
            st.success(
                f"Готово: файлов {ok}, строк {total_rows}, "
                f"удалено строк с нулевой компенсацией {total_skipped}."
            )
            st.download_button(
                "Скачать готовые шаблоны ZIP",
                data=zip_buffer.getvalue(),
                file_name=f"Готовые шаблоны_паспорт_{passport.strip()}.zip",
                mime="application/zip",
                use_container_width=True,
            )

        if bad:
            st.error(f"Не обработано файлов: {bad}")

        with st.expander("Подробности обработки", expanded=bool(bad)):
            for status, name, kept, skipped, error in results:
                if error:
                    st.write(f"{status} **{name}** — {error}")
                else:
                    st.write(f"{status} **{name}** — {kept} строк; удалено нулевых: {skipped}")
